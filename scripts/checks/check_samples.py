#!/usr/bin/env python3
"""Sample check: every sampled transaction exists, unchanged, in the cleaned data and in the council's workbook.

For every row of outputs/samples/*.csv, and every row ID quoted in
audit-program.md, the row ID (file|sheet|Excel row):
  - is in analytics-full/spend_clean.csv, and a sample row has the same
    supplier and amount there;
  - is at that Excel row of that sheet in the council's workbook in
    data/spend/ (opened read-only), with the same supplier name and net
    amount.
The workbook's header row is found by its column names, not assumed.
A sample file with no rows, or with a duplicate row ID, also fails.

Usage: check_samples.py [--out-dir outputs] [--full-dir …] [--spend-dir data/spend]
Exit 1 if the check fails.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import Result, add_path_args, paths, read_text  # noqa: E402

ROW_ID_RE = re.compile(r"([^\s|`]+\.xlsx)\|([^|\n`]+?)\|(\d+)")
SUPPLIER_COL, AMOUNT_COL = "Supplier name", "Net amount"
HEADER_SCAN_ROWS = 30


def same_supplier(a, b) -> bool:
    return re.sub(r"\s+", " ", str(a or "")).strip() == re.sub(r"\s+", " ", str(b or "")).strip()


def same_amount(a, b) -> bool:
    try:
        return abs(float(a) - float(b)) < 0.005
    except (TypeError, ValueError):
        return False


def workbook_rows(spend_dir: Path, wanted: dict, res: Result) -> dict:
    """{row_id: (supplier, amount)} read from the workbooks, one pass per sheet."""
    from openpyxl import load_workbook
    found = {}
    for (fname, sheet), rows in wanted.items():
        path = spend_dir / fname
        if not path.is_file():
            res.fail(f"{fname}: workbook not found in {spend_dir}")
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet not in wb.sheetnames:
                res.fail(f"{fname}: no sheet '{sheet}'")
                continue
            ws = wb[sheet]
            cols, last = None, max(rows)
            for n, vals in enumerate(ws.iter_rows(values_only=True), 1):
                if cols is None:
                    names = [str(v).strip() if v is not None else "" for v in vals]
                    if SUPPLIER_COL in names and AMOUNT_COL in names:
                        cols = (names.index(SUPPLIER_COL), names.index(AMOUNT_COL))
                    elif n >= HEADER_SCAN_ROWS:
                        res.fail(f"{fname}: no header row with '{SUPPLIER_COL}' and '{AMOUNT_COL}' "
                                 f"in the first {HEADER_SCAN_ROWS} rows")
                        break
                    continue
                if n in rows:
                    found[f"{fname}|{sheet}|{n}"] = (vals[cols[0]] if cols[0] < len(vals) else None,
                                                     vals[cols[1]] if cols[1] < len(vals) else None)
                if n >= last:
                    break
        finally:
            wb.close()
    return found


def run(p: dict) -> Result:
    res = Result("samples", what="transactions")
    if not p["spend_clean"].is_file():
        res.fail(f"{p['spend_clean']}: missing")
        return res
    clean = {}
    with p["spend_clean"].open(newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            clean[row["row_id"]] = (row["supplier_raw"], row["amount"])

    # row IDs to check: (where, row_id, supplier or None, amount or None)
    items = []
    if not p["samples"].is_dir():
        res.fail(f"{p['samples']}: missing")
    else:
        files = sorted(p["samples"].glob("*.csv"))
        if not files:
            res.fail(f"{p['samples']}: no sample files")
        for f in files:
            seen = set()
            with f.open(newline="", encoding="utf-8") as fh:
                rows = list(csv.DictReader(fh))
            if not rows:
                res.fail(f"samples/{f.name}: no rows")
            for i, row in enumerate(rows, 2):
                rid = row.get("row_id", "")
                if rid in seen:
                    res.fail(f"samples/{f.name}:{i}: row {rid} appears twice")
                seen.add(rid)
                items.append((f"samples/{f.name}:{i}", rid, row.get("supplier"), row.get("amount")))
    prog = read_text(p["audit-program.md"], res)
    n_prog = 0
    for n, line in enumerate((prog or "").splitlines(), 1):
        for m in ROW_ID_RE.finditer(line):
            items.append((f"audit-program.md:{n}", m.group(0), None, None))
            n_prog += 1

    # 1. spend_clean.csv
    wanted = defaultdict(set)
    for where, rid, sup, amt in items:
        res.checked += 1
        m = ROW_ID_RE.fullmatch(rid or "")
        if not m:
            res.fail(f"{where}: '{rid}' is not a row ID (file|sheet|row)")
            continue
        wanted[(m.group(1), m.group(2))].add(int(m.group(3)))
        if rid not in clean:
            res.fail(f"{where}: {rid} is not in spend_clean.csv")
            continue
        c_sup, c_amt = clean[rid]
        if sup is not None and not same_supplier(sup, c_sup):
            res.fail(f"{where}: {rid} supplier differs from spend_clean.csv")
        if amt is not None and not same_amount(amt, c_amt):
            res.fail(f"{where}: {rid} amount {amt} differs from spend_clean.csv ({c_amt})")

    # 2. the council's workbook, at the sheet row
    wb = workbook_rows(p["spend_dir"], wanted, res)
    for where, rid, sup, amt in items:
        if not ROW_ID_RE.fullmatch(rid or ""):
            continue
        if rid not in wb:
            res.fail(f"{where}: {rid} not found at that row of the workbook")
            continue
        w_sup, w_amt = wb[rid]
        sup = sup if sup is not None else clean.get(rid, (None, None))[0]
        amt = amt if amt is not None else clean.get(rid, (None, None))[1]
        if sup is not None and not same_supplier(sup, w_sup):
            res.fail(f"{where}: {rid} supplier differs from the workbook row")
        if amt is not None and not same_amount(amt, w_amt):
            res.fail(f"{where}: {rid} amount {amt} differs from the workbook row ({w_amt})")
    res.notes.append(f"{len(items) - n_prog} sampled transaction(s) in {len({w.split(':')[0] for w, *_ in items if w.startswith('samples/')})} "
                     f"file(s); {n_prog} row ID(s) quoted in audit-program.md; "
                     f"{sum(len(v) for v in wanted.values())} workbook row(s) read")
    return res


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    add_path_args(ap)
    res = run(paths(ap.parse_args()))
    for f in res.failures:
        print(f"FAIL {f}")
    for n in res.notes:
        print(f"note {n}")
    print(f"samples: {res.status} ({res.checked} {res.what} checked, {len(res.failures)} failure(s))")
    return 1 if res.failures else 0


if __name__ == "__main__":
    sys.exit(main())
